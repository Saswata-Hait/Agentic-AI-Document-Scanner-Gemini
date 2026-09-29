from crewai.tools import tool

import os
import pandas as pd

from pypdf import PdfReader
from docx import Document


# ============================================================
# PDF Reader Tool
# ============================================================

@tool
def readPDF(file_path: str) -> str:
    """
    Read text content from a PDF file.
    """

    if not os.path.exists(file_path):
        return f"File not found: {file_path}"

    try:

        reader = PdfReader(file_path)

        text = ""

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            page_text = page.extract_text()

            if page_text:

                text += (
                    f"\n[PAGE {page_number}]\n"
                )

                text += page_text
                text += "\n"

        return text

    except Exception as e:

        return f"Error reading PDF: {str(e)}"


# ============================================================
# CSV Reader Tool
# ============================================================

@tool
def readCSV(file_path: str) -> str:
    """
    Read CSV file and return its content as text.
    """

    if not os.path.exists(file_path):
        return f"File not found: {file_path}"

    try:

        df = pd.read_csv(file_path)

        return df.to_string(
            index=False
        )

    except Exception as e:

        return f"Error reading CSV: {str(e)}"


# ============================================================
# Excel Reader Tool
# ============================================================

@tool
def readExcel(file_path: str) -> str:
    """
    Read Excel file and all worksheets
    and return their content as text.
    """

    if not os.path.exists(file_path):
        return f"File not found: {file_path}"

    try:

        excel_file = pd.ExcelFile(
            file_path
        )

        result = ""

        for sheet in excel_file.sheet_names:

            result += (
                "\n\n"
                "========================================\n"
                f"SHEET: {sheet}\n"
                "========================================\n"
            )

            df = pd.read_excel(
                file_path,
                sheet_name=sheet
            )

            result += df.to_string(
                index=False
            )

            result += "\n"

        return result

    except Exception as e:

        return f"Error reading Excel: {str(e)}"


# ============================================================
# Word Reader Tool
# ============================================================

@tool
def readWord(file_path: str) -> str:
    """
    Read DOCX file and return its text content.
    """

    if not os.path.exists(file_path):
        return f"File not found: {file_path}"

    try:

        document = Document(
            file_path
        )

        text = ""

        # --------------------------------------------
        # Read paragraphs
        # --------------------------------------------

        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                text += (
                    paragraph.text
                    + "\n"
                )


        # --------------------------------------------
        # Read tables
        # --------------------------------------------

        for table_number, table in enumerate(
            document.tables,
            start=1
        ):

            text += (
                "\n\n"
                f"[TABLE {table_number}]\n"
            )

            for row in table.rows:

                row_data = []

                for cell in row.cells:

                    row_data.append(
                        cell.text.strip()
                    )

                text += (
                    " | ".join(row_data)
                    + "\n"
                )

        return text

    except Exception as e:

        return f"Error reading DOCX: {str(e)}"